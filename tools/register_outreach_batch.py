"""Register an existing-contact XLSX as an auditable batch. Never sends email."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from sales_automation.config import load_config
from sales_automation.db import Database, Repository
from sales_automation.services.outreach_batches import OutreachBatchService


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--config", default="config.yaml")
    parser.add_argument("--owner", required=True)
    parser.add_argument("--name", required=True)
    parser.add_argument("--source-name", default="")
    parser.add_argument("--claim-public", action="store_true")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    config = load_config(args.config)
    db = Database(config)
    db.bind_system_actor()
    repo = Repository(db)
    with db.connect() as conn:
        owner = conn.execute(
            "SELECT id,username,display_name,role,active FROM sales_users WHERE lower(username)=lower(%s) AND active",
            (args.owner,),
        ).fetchone()
    if not owner or owner["role"] != "sales":
        raise ValueError("Select an active salesperson")
    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)
    try:
        records = workbook.worksheets[0].iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(records)]
        id_column = next((i for i, value in enumerate(headers)
                          if value in {"contact_id", "\u5ba2\u6237ID"}), None)
        if id_column is None:
            raise ValueError("A contact_id column is required; new contacts must use normal import first")
        original = {}
        for values in records:
            value = values[id_column]
            if value is None:
                continue
            contact_id = int(value)
            if contact_id <= 0 or str(contact_id) in original:
                raise ValueError("Invalid or duplicate contact_id")
            original[str(contact_id)] = dict(zip(headers, values))
    finally:
        workbook.close()
    ids = [int(value) for value in original]
    if not 1 <= len(ids) <= 1000:
        raise ValueError("Batch must contain 1-1000 contacts")
    with db.connect() as conn:
        contacts = conn.execute(
            "SELECT id,pool_type,owner_user_id,sabcd_stage FROM contacts WHERE id=ANY(%s)", (ids,)
        ).fetchall()
    own = [row["id"] for row in contacts if row["pool_type"] == "private" and row["owner_user_id"] == owner["id"]]
    public = [row["id"] for row in contacts if row["pool_type"] == "public" and row["sabcd_stage"] != "S"]
    eligible = set(own + (public if args.claim_public else []))
    excluded = [value for value in ids if value not in eligible]
    report = {"input": len(ids), "owned": len(own), "claimable": len(public),
              "excluded": len(excluded), "excluded_ids": excluded, "applied": False, "sent": 0}
    if args.apply:
        source_name = args.source_name or args.xlsx.name
        source_hash = hashlib.sha256(args.xlsx.read_bytes()).hexdigest()
        if args.claim_public:
            repo.assign_public_contacts_to_owner(public, owner_user_id=owner["id"],
                owner_name=owner["display_name"], assignment_source="outreach_batch")
        result = OutreachBatchService(repo).create(user=owner, name=args.name,
            contact_ids=[value for value in ids if value in eligible], source_ref=source_name,
            source_rows=original, metadata={"source_sha256": source_hash,
                "input_count": len(ids), "excluded_count": len(excluded),
                "preparation_status": "awaiting_research"})
        campaign_id = result["batch"]["id"]
        repo.record_audit_log(user=owner, action="register_outreach_batch", target_type="campaign",
            target_id=campaign_id, summary="Register existing contacts for reviewed outreach",
            metadata={"input_count": len(ids), "registered": result["total"],
                      "excluded_count": len(excluded), "source_sha256": source_hash})
        report.update(applied=True, campaign_id=campaign_id, registered=result["total"])
    print(json.dumps(report, ensure_ascii=True))


if __name__ == "__main__":
    main()
